from collections import defaultdict
from pathlib import Path
import openpyxl

DATA_DIR = Path(__file__).resolve().parent / 'data'
OUTPUT = Path(__file__).resolve().parent / 'conteggio_PI_per_universita.md'

# Raggruppa per codice fiscale dell'ente, quando disponibile,
# evitando di separare la stessa universita per varianti del nome.
enti = defaultdict(lambda: {'nomi': set(), 'persone': set(), 'progetti': set()})
file_letti = 0

for file in sorted(DATA_DIR.glob('*.xlsx')):
    if file.name.startswith('~$'):
        continue
    wb = openpyxl.load_workbook(file, read_only=True, data_only=True)
    if 'Progetti' not in wb.sheetnames:
        print(f'Salto {file.name}: foglio Progetti assente')
        wb.close()
        continue
    ws = wb['Progetti']
    righe = ws.iter_rows(values_only=True)
    intestazioni = None
    for riga in righe:
        normalizzate = [str(x).strip().casefold() if x is not None else '' for x in riga]
        if 'ruolo' in normalizzate and 'ateneo/ente' in normalizzate:
            intestazioni = normalizzate
            break
    if intestazioni is None:
        print(f'Salto {file.name}: intestazioni non trovate')
        wb.close()
        continue

    idx_ruolo = intestazioni.index('ruolo')
    idx_ente = intestazioni.index('ateneo/ente')
    idx_persona = next(i for i, x in enumerate(intestazioni) if 'responsabile unit' in x)
    idx_cf = next((i for i, x in enumerate(intestazioni) if 'codice fiscale' in x), None)
    idx_progetto = intestazioni.index('codice progetto') if 'codice progetto' in intestazioni else None

    for riga in righe:
        if str(riga[idx_ruolo] or '').strip().upper() != 'PI':
            continue
        ente = str(riga[idx_ente] or '').strip()
        persona = str(riga[idx_persona] or '').strip()
        cf = str(riga[idx_cf] or '').strip() if idx_cf is not None else ''
        if not ente or not persona:
            continue
        chiave = cf if cf else ente.casefold()
        enti[chiave]['nomi'].add(ente)
        enti[chiave]['persone'].add(persona.casefold())
        if idx_progetto is not None and riga[idx_progetto]:
            enti[chiave]['progetti'].add(str(riga[idx_progetto]).strip())
    file_letti += 1
    wb.close()

risultati = sorted(
    [(sorted(d['nomi'], key=lambda x: (len(x), x))[0], len(d['persone']), len(d['progetti']))
     for d in enti.values()],
    key=lambda x: (-x[1], x[0].casefold())
)
righe_md = [
    '# Numero di persone PI per università/ente',
    '',
    f'File Excel analizzati: **{file_letti}**.',
    '',
    '| Università / Ente | Persone PI distinte | Progetti con PI |',
    '|---|---:|---:|',
]
for ente, n_persone, n_progetti in risultati:
    righe_md.append(f'| {ente.replace("|", "/")} | {n_persone} | {n_progetti} |')
righe_md += [
    '',
    f'**Totale persone PI (somma per ente): {sum(x[1] for x in risultati)}**',
    '',
    'Nota: sono considerate solo le righe con ruolo `PI`, non quelle `RU`. '
    'Il conteggio delle persone distinte usa il nome del responsabile all’interno di ogni ente; '
    'gli enti sono raggruppati per codice fiscale, ove presente.',
    ''
]
OUTPUT.write_text('\n'.join(righe_md), encoding='utf-8')
print(f'Analizzati {file_letti} file; trovati {len(risultati)} enti; creato: {OUTPUT}')
print('\n'.join(righe_md[:12]))
