#!/usr/bin/env python3
"""Aggrega il contributo MUR assegnato agli atenei/enti nei file PRIN Excel.

Esempio: python finanziamenti_per_ateneo.py data --md riepilogo_atenei.md
Output Markdown sia sul terminale sia nel file .md; CSV opzionale con --csv.
Dipendenza: pip install openpyxl
"""
import argparse
import csv
import sys
from collections import defaultdict
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import load_workbook


def euro(valore):
    """Converte numeri Excel o importi italiani (es. 1.234,56) in Decimal."""
    if valore is None or str(valore).strip() == "":
        return Decimal(0)
    if isinstance(valore, (int, float, Decimal)):
        return Decimal(str(valore))
    testo = str(valore).replace("€", "").replace(" ", "").strip()
    if "," in testo:
        testo = testo.replace(".", "").replace(",", ".")
    try:
        return Decimal(testo)
    except InvalidOperation as exc:
        raise ValueError(f"Importo non valido: {valore!r}") from exc


def formato_euro(numero):
    return f"{numero:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")


def calcola(cartella):
    importi = defaultdict(lambda: Decimal(0))
    nomi = {}
    visti = set()
    letti = 0

    for file in sorted(cartella.glob("*.xlsx")):
        if file.name.startswith("~$"):
            continue
        wb = load_workbook(file, read_only=True, data_only=True)
        try:
            if "Unita" not in wb.sheetnames:
                print(f"File ignorato (manca il foglio 'Unita'): {file.name}")
                continue
            righe = wb["Unita"].iter_rows(values_only=True)
            intestazioni = [str(c).strip() if c is not None else "" for c in next(righe)]
            colonne = {titolo: idx for idx, titolo in enumerate(intestazioni)}
            richieste = ["Ateneo/Ente", "Codice fiscale", "MUR assegnato (€)", "Codice unità"]
            mancanti = [titolo for titolo in richieste if titolo not in colonne]
            if mancanti:
                raise ValueError(f"{file.name}: colonne mancanti: {', '.join(mancanti)}")

            totale_file = Decimal(0)
            for riga in righe:
                ateneo = str(riga[colonne["Ateneo/Ente"]] or "").strip()
                cf = str(riga[colonne["Codice fiscale"]] or "").strip()
                unita = str(riga[colonne["Codice unità"]] or "").strip()
                if not unita:
                    raise ValueError(f"{file.name}: riga senza codice unità")
                if not ateneo:
                    raise ValueError(f"{file.name}: ateneo/ente assente per {unita}")
                # Usa il codice fiscale per riunire varianti del nome dello stesso ente.
                chiave = cf or ateneo.casefold()
                chiave_unita = (chiave, unita)
                if chiave_unita in visti:
                    continue
                visti.add(chiave_unita)
                nomi.setdefault(chiave, ateneo)
                importo = euro(riga[colonne["MUR assegnato (€)"]])
                importi[chiave] += importo
                totale_file += importo
            # Verifica che la somma delle unità sia uguale al totale dei progetti.
            if "Progetti" in wb.sheetnames:
                pr = wb["Progetti"].iter_rows(values_only=True)
                headers = [str(c).strip() if c is not None else "" for c in next(pr)]
                if "Contributo MUR assegnato (€)" in headers:
                    pos = headers.index("Contributo MUR assegnato (€)")
                    totale_atteso = sum((euro(r[pos]) for r in pr), Decimal(0))
                    if totale_file != totale_atteso:
                        raise ValueError(f"{file.name}: totale Unita {formato_euro(totale_file)} != "
                                         f"totale Progetti {formato_euro(totale_atteso)}. "
                                         "Verificare righe mancanti nella conversione.")
            letti += 1
        finally:
            wb.close()

    ordinati = sorted(importi.items(), key=lambda item: (-item[1], nomi[item[0]]))
    totale = sum(importi.values(), Decimal(0))
    return [(nomi[k], k, somma, somma / totale * 100 if totale else Decimal(0))
            for k, somma in ordinati], totale, letti


def crea_markdown(risultati, totale, nfile):
    """Restituisce un documento Markdown completo con la classifica."""
    righe = [
        "# Finanziamenti PRIN 2026 per università/ente",
        "",
        f"**File Excel analizzati:** {nfile}  ",
        f"**Enti distinti:** {len(risultati)}  ",
        f"**Finanziamento MUR complessivo:** {formato_euro(totale)} €",
        "",
        "| Pos. | Università / Ente | Totale finanziato (EUR) | % sul totale |",
        "|---:|---|---:|---:|",
    ]
    for posizione, (nome, _, importo, quota) in enumerate(risultati, 1):
        nome_md = nome.replace("|", "\\|").replace("\n", " ").replace("\r", " ")
        percentuale = f"{quota:.2f}".replace(".", ",")
        righe.append(f"| {posizione} | {nome_md} | {formato_euro(importo)} € | {percentuale}% |")
    righe.append(f"| | **TOTALE** | **{formato_euro(totale)} €** | **100,00%** |")
    righe.extend(["", "*La percentuale è calcolata rispetto alla somma dei contributi MUR assegnati a tutte le unità dei file analizzati.*", ""])
    return "\n".join(righe)


def main():
    parser = argparse.ArgumentParser(description="Totale finanziamenti MUR e quota percentuale per ateneo/ente")
    parser.add_argument("cartella", nargs="?", default="data", type=Path,
                        help="Cartella dei file XLSX (predefinito: data)")
    parser.add_argument("--csv", type=Path, help="Esporta il riepilogo in CSV")
    parser.add_argument("--md", type=Path, default=Path("riepilogo_finanziamenti_atenei.md"),
                        help="Percorso del file Markdown (predefinito: riepilogo_finanziamenti_atenei.md)")
    args = parser.parse_args()
    if not args.cartella.is_dir():
        parser.error(f"Cartella non trovata: {args.cartella}")

    risultati, totale, nfile = calcola(args.cartella)
    if nfile == 0:
        parser.error("Nessun file con foglio 'Unita' trovato")

    markdown = crea_markdown(risultati, totale, nfile)
    print(markdown, end="")
    args.md.write_text(markdown, encoding="utf-8")
    print(f"File Markdown scritto in: {args.md}", file=sys.stderr)

    if args.csv:
        with args.csv.open("w", newline="", encoding="utf-8-sig") as f:
            writer = csv.writer(f, delimiter=";")
            writer.writerow(["Ateneo/Ente", "Codice fiscale", "Totale MUR assegnato (€)", "Percentuale (%)"])
            for nome, cf, somma, quota in risultati:
                writer.writerow([nome, cf, str(somma), f"{quota:.4f}"])
        print(f"CSV scritto in: {args.csv}", file=sys.stderr)


if __name__ == "__main__":
    main()
