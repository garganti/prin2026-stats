#!/usr/bin/env python3
"""Importa le sette nuove Tabelle B dai PDF in data/*.xlsx.

Eseguire dalla root: python code/importa_nuovi_settori.py
Per ogni progetto confronta la somma degli importi delle unita' con il totale
stampato sul PDF. Richiede pdftotext (poppler) e openpyxl.
I sette file Excel vengono salvati in data/; gli script di analisi producono solo Markdown.
"""
from __future__ import annotations

import re
import subprocess
from collections import defaultdict
from copy import copy
from decimal import Decimal
from pathlib import Path
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
TEMPLATE = ROOT / 'data' / 'HPC_Tabella_B_Progetti_finanziati.xlsx'
SECTORS = {'HT': (13, 18170142), 'SH3': (7, 8023944), 'SH8': (8, 8873946),
           'LS8': (8, 8317931), 'PE4': (6, 6806601), 'LS9': (14, 14827842),
           'LS5': (8, 9393088)}
AMOUNT = r'(\d[\d.]*,\d{2})'
PROJECT_RE = re.compile(r'^\s*(\d+)\.\s+(2026[A-Z0-9]+)\s+-\s+(.+?)\s*$')
UNIT_RE = re.compile(r'^\s*(\d+)\s+(PI|RU)\s+(.*?)\s+(2026[A-Z0-9]+_\d{3})\s+'
    r'(.*?)\s*(\d{11})\s+' + AMOUNT + r'\s+' + AMOUNT + r'\s+' + AMOUNT + r'\s+' + AMOUNT + r'\s*$')
TOTAL_RE = re.compile(r'^\s*Totale\s+' + AMOUNT + r'\s+' + AMOUNT + r'\s+' + AMOUNT + r'\s*$')
GRAND_RE = re.compile(r'^.*TOTALE COMPLESSIVO\s+' + AMOUNT + r'\s*$')
# In these cases the institution name is on lines above/below its financial row.
# They are transcribed from the source PDF. Others are recovered by matching
# the institution fiscal identifier in the existing dataset.
NAMES_BY_CF = {
    '80024610158': 'Università Commerciale "Luigi Bocconi" MILANO',
    '95003210655': 'CONSERVATORIO DI MUSICA DI SALERNO',
    '80010280347': 'CONSERVATORIO DI MUSICA DI PARMA',
    '97394340588': 'Università Telematica Internazionale UNINETTUNO',
}


def number(text):
    d = Decimal(text.replace('.', '').replace(',', '.'))
    if d != d.to_integral_value():
        raise ValueError(f'Importo non intero nel PDF: {text}')
    return int(d)


def institution_catalogue():
    names = defaultdict(set)
    for path in sorted((ROOT / 'data').glob('*.xlsx')):
        if path.stem.split('_')[0] in SECTORS:
            continue
        w = load_workbook(path, read_only=True, data_only=True)
        try:
            ws = w['Unita'] if 'Unita' in w else w['Progetti']
            ws.reset_dimensions()
            for row in ws.iter_rows(values_only=True):
                if len(row) > 7:
                    cf = str(row[7] or '').strip()
                    nm = str(row[6] or '').strip()
                    if len(cf) == 11 and nm and not nm.lower().startswith(('ateneo', 'codice')):
                        names[cf].add(nm)
        finally:
            w.close()
    return {cf: sorted(v, key=lambda s:(len(s),s))[0] for cf,v in names.items()}


def read_pdf(sector, catalog):
    matches = list(ROOT.glob(f'{sector}_Tabella_B_*.pdf'))
    if not matches:
        matches = list((ROOT / 'fonti').glob(f'{sector}_Tabella_B_*.pdf'))
    if len(matches) != 1:
        raise ValueError(f'{sector}: atteso un PDF nella root, trovati: {matches}')
    pdf=matches[0]
    raw = subprocess.check_output(['pdftotext', '-layout', str(pdf), '-'], text=True)
    projects, units, totals = [], [], []
    grand = None
    title_open = False
    for rawline in raw.splitlines():
        stripped = ' '.join(rawline.split())
        if m := PROJECT_RE.match(rawline):
            n, code, title = m.groups()
            projects.append({'n': int(n), 'code': code, 'title': title.strip()})
            title_open = True
            continue
        if m := UNIT_RE.match(stripped):
            title_open = False
            n, role, person, unitcode, name, cf, cof, mur, cost, assigned = m.groups()
            if not name:
                name = NAMES_BY_CF.get(cf) or catalog.get(cf) or ''
            if not name:
                raise ValueError(f'{sector}: ente non trovato per {unitcode} ({cf})')
            u = dict(n=int(n), role=role, person=person.strip(), id=unitcode,
                     project=unitcode.split('_')[0], name=name, cf=cf,
                     cof=number(cof), mur=number(mur), cost=number(cost),
                     assigned=number(assigned))
            if u['cof']+u['mur'] != u['cost'] or u['mur']!=u['assigned']:
                raise ValueError(f'{unitcode}: importi di unita discordanti')
            units.append(u)
            continue
        if m := TOTAL_RE.match(stripped):
            title_open = False
            totals.append(tuple(map(number,m.groups())))
            continue
        if m := GRAND_RE.match(stripped):
            title_open = False
            grand = number(m.group(1))
            continue
        # Some long project titles occupy the next physical line.
        # Only accept nonempty lines before any unit has been seen.
        if title_open and stripped:
            if stripped.startswith(('n.b.', 'TOTALE COMPLESSIVO')) or ' - Tabella B - ' in stripped:
                title_open=False
            elif (len(rawline) - len(rawline.lstrip()) < 18 and
                  not re.match(r'^\d+\s+(PI|RU)\b', stripped)):
                projects[-1]['title'] += ' ' + stripped
    expected_count, expected_grand = SECTORS[sector]
    if len(projects) != expected_count or len(totals) != expected_count or grand != expected_grand:
        raise ValueError(f'{sector}: parse progetti {len(projects)}, totali {len(totals)}, '
                         f'MUR={grand}; attesi {expected_count}, {expected_grand}')
    if len({u['id'] for u in units})!=len(units):
        raise ValueError(f'{sector}: codici unita duplicati')
    grouped = defaultdict(list)
    for u in units:
        grouped[u['project']].append(u)
    if set(grouped) != {p['code'] for p in projects}:
        raise ValueError(f'{sector}: progetti senza dettaglio o unita sconosciute')
    for proj, (cof,mur,assigned) in zip(projects,totals):
        group = sorted(grouped[proj['code']],key=lambda u:u['n'])
        if sum(u['role']=='PI' for u in group)!=1:
            raise ValueError(f'{proj["code"]}: PI non unico')
        if [u['n'] for u in group]!=list(range(1,len(group)+1)):
            raise ValueError(f'{proj["code"]}: numerazione non progressiva')
        if (sum(u['cof'] for u in group), sum(u['mur'] for u in group))!=(cof,mur) or mur!=assigned:
            raise ValueError(f'{proj["code"]}: subtotali diversi dal PDF')
        proj.update(cof=cof, mur=mur, n_units=len(group))
    if sum(p['mur'] for p in projects)!=grand:
        raise ValueError(f'{sector}: totale finanziamenti non riconciliato')
    units.sort(key=lambda u:(next(p['n'] for p in projects if p['code']==u['project']),u['n']))
    return pdf, projects, units


def create(sector, catalogue):
    pdf,projects,units = read_pdf(sector,catalogue)
    w=load_workbook(TEMPLATE)
    a,b=w['Riepilogo'],w['Progetti']
    style_a=[copy(a.cell(5,j)._style) for j in range(1,7)]
    style_b=[copy(b.cell(5,j)._style) for j in range(1,14)]
    a.delete_rows(5,a.max_row-4)
    b.delete_rows(5,b.max_row-4)
    a['A1']=f'{sector} - Tabella B - Progetti finanziati'
    b['A1']=f'{sector} - Tabella B - Progetti finanziati - Dettaglio delle unità'
    b['A2']='Dati trascritti dal PDF; stato PI under 40 non verificato'
    for i,p in enumerate(projects,5):
        for j,v in enumerate([p['n'],p['code'],p['title'],p['n_units'],p['mur'],p['cof']],1):
            c=a.cell(i,j,v);c._style=copy(style_a[j-1])
    final=len(projects)+5
    for j,v in [(1,'TOTALE'),(4,len(units)),(5,sum(p['mur'] for p in projects)),(6,sum(p['cof'] for p in projects))]:
        c=a.cell(final,j,v);c._style=copy(style_a[j-1]);f=copy(c.font);f.bold=True;c.font=f
    a.cell(final+2,1,'Fonte: '+pdf.name)
    mapping={p['code']:p for p in projects}
    for i,u in enumerate(units,5):
        p=mapping[u['project']]
        values=[p['n'],p['code'],p['title'],u['role'],u['person'],u['id'],u['name'],u['cf'],
                u['cof'],u['mur'],u['cost'],u['assigned'],'Non verificato']
        for j,v in enumerate(values,1):
            c=b.cell(i,j,v);c._style=copy(style_b[j-1])
        b.cell(i,8).number_format='@'
    a.auto_filter.ref=f'A4:F{len(projects)+4}'
    b.auto_filter.ref=f'A4:M{len(units)+4}'
    dst=ROOT/'data'/f'{sector}_Tabella_B_Progetti_finanziati.xlsx'
    w.save(dst);w.close()
    # Check values were written and kept intact.
    check=load_workbook(dst,read_only=True,data_only=True)
    try:
        assert sum(1 for row in check['Progetti'].iter_rows(min_row=5,values_only=True) if row[1])==len(units)
        assert sum(row[4] for row in check['Riepilogo'].iter_rows(min_row=5,max_row=4+len(projects),values_only=True))==sum(p['mur'] for p in projects)
    finally:
        check.close()
    print(f'{sector}: {len(projects)} progetti, {len(units)} unita, {len(units)-len(projects)} RU, MUR {sum(p["mur"] for p in projects):,} EUR')
    return dst


if __name__=='__main__':
    catalogue=institution_catalogue()
    for sector in SECTORS:
        create(sector,catalogue)
