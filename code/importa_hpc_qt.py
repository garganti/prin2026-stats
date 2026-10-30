#!/usr/bin/env python3
"""Converte le Tabelle B HPC e QT da PDF in Excel PRIN 2026 validati.

Uso: python code/importa_hpc_qt.py [HPC | QT | ALL]
Dipendenze: openpyxl, poppler (`pdftotext`).
I file .xlsx risultanti sono salvati in data/ senza modificare gli altri settori.

Le cinque righe multilinea sono ricostruite esplicitamente sulla base dei PDF.
Ogni progetto e ogni contributo sono verificati rispetto ai totali del PDF.
"""
from __future__ import annotations

import re
import subprocess
import sys
from collections import defaultdict
from copy import copy
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import PatternFill

ROOT = Path(__file__).resolve().parent.parent
TEMPLATE = ROOT / 'data' / 'LS6_Tabella_B_Progetti_finanziati.xlsx'
AMOUNT = r'(\d[\d.]*,\d{2})'
PROJECT_RE = re.compile(r'^\s*(\d+)\.\s+(2026[A-Z0-9]+)\s+-\s+(.+?)\s*$')
UNIT_RE = re.compile(
    r'^\s*(\d+)\s+(PI|RU)\s+(.*?)\s+(2026[A-Z0-9]+_\d{3})\s+'
    r'(.*?)\s+(\d{11})\s+' + AMOUNT + r'\s+' + AMOUNT + r'\s+' + AMOUNT + r'\s+' + AMOUNT + r'\s*$'
)
TOTAL_RE = re.compile(r'^\s*Totale\s+' + AMOUNT + r'\s+' + AMOUNT + r'\s+' + AMOUNT + r'\s*$')
GRAND_RE = re.compile(r'^.*TOTALE COMPLESSIVO\s+' + AMOUNT + r'\s*$')

# Righe in cui la denominazione o il responsabile sono distribuiti su piu' righe.
# I valori sono trascritti letteralmente dalle tabelle allegate.
MULTILINE = {
    '2026H879TM_003': ('RU', 'PANZANELLI FRATONI Maria Alessandra',
                       'Università degli Studi di TORINO', '80088230018', 80225, 291018, 371243),
    '2026HPEYTW_005': ('RU', 'MULARGIA Simone',
                       'Libera Università degli Studi "Maria SS.Assunta" - LUMSA',
                       '02635620582', 64875, 278000, 342875),
    '2026HRKPT7_004': ('RU', 'NIEDDU Luciano',
                       'Università degli Studi Internazionali di ROMA (UNINT)',
                       '97136680580', 106000, 169839, 275839),
    '2026HL78Z9_005': ('RU', 'COZZINI Stefano',
                       'Area di Ricerca Scientifica e Tecnologica di Trieste - Area Science Park',
                       '00531590321', 79640, 234669, 314309),
    '2026HZ9RR3_004': ('RU', 'AULETTA Gennaro',
                       'Università degli Studi di CASSINO e del LAZIO MERIDIONALE',
                       '81006500607', 72875, 167700, 240575),
}
SETTINGS = {
    'HPC': {'n_projects': 4, 'n_units': 20, 'grand': 5541856, 'green': set()},
    'QT': {'n_projects': 5, 'n_units': 27, 'grand': 6618858, 'green': {'2026HS3N55'}},
}


def amount(s: str) -> int:
    value = Decimal(s.replace('.', '').replace(',', '.'))
    if value != value.to_integral_value():
        raise ValueError(f'Importo non intero: {s}')
    return int(value)


def parse_pdf(sector: str):
    pdf = ROOT / f'{sector}_Tabella_B_Progetti finanziati.pdf'
    if not pdf.is_file():
        pdf = ROOT / 'fonti' / pdf.name
    if not pdf.is_file():
        raise FileNotFoundError(pdf)
    raw = subprocess.check_output(['pdftotext', '-layout', str(pdf), '-'], text=True)
    projects, units, totals = [], [], []
    grand_total = None
    continuation = False
    for line in raw.splitlines():
        normalized = ' '.join(line.split())
        if match := PROJECT_RE.match(line):
            n, code, title = match.groups()
            projects.append({'n': int(n), 'codice': code, 'titolo': title.strip()})
            continuation = True
            continue
        if continuation and normalized and not re.match(r'^\d+\s+(PI|RU)\s', normalized):
            # Only the long first HPC title wraps onto another physical line.
            if sector == 'HPC' and len(projects) == 1 and normalized.startswith('the bibliographical analysis'):
                projects[-1]['titolo'] += ' ' + normalized
                continue
        if normalized and re.match(r'^\d+\s+(PI|RU)\s', normalized):
            continuation = False
        if match := UNIT_RE.match(normalized):
            number, role, person, code, institution, cf, cofin, mur, cost, assigned = match.groups()
            record = dict(n=int(number), ruolo=role, responsabile=person.strip(),
                          codice=code, progetto=code.split('_')[0], ente=institution.strip(),
                          cf=cf, cof=amount(cofin), mur=amount(mur),
                          costo=amount(cost), assegnato=amount(assigned))
            units.append(record)
        elif match := TOTAL_RE.match(normalized):
            totals.append(tuple(amount(t) for t in match.groups()))
        elif match := GRAND_RE.match(normalized):
            grand_total = amount(match.group(1))

    # Add the PDF rows whose name/institution line wraps independently of amounts.
    codes_found = {u['codice'] for u in units}
    for code, (role, person, institution, cf, cofin, mur, cost) in MULTILINE.items():
        if code not in raw:
            continue
        if code in codes_found:
            raise ValueError(f'{code} parsed automatically and hardcoded: remove manual override')
        unit_n = int(code.split('_')[-1])
        units.append(dict(n=unit_n, ruolo=role, responsabile=person, codice=code,
                          progetto=code.split('_')[0], ente=institution, cf=cf,
                          cof=cofin, mur=mur, costo=cost, assegnato=mur))

    setting = SETTINGS[sector]
    if (len(projects), len(units), len(totals), grand_total) != (
        setting['n_projects'], setting['n_units'], setting['n_projects'], setting['grand']
    ):
        raise ValueError(f'{sector}: estrazione incompleta: {len(projects)} progetti, '
                         f'{len(units)} unita, {len(totals)} subtotali, finale {grand_total}')
    by_project = defaultdict(list)
    codes = set()
    for u in units:
        if u['codice'] in codes:
            raise ValueError('Codice unità duplicato: ' + u['codice'])
        codes.add(u['codice'])
        if u['costo'] != u['cof'] + u['mur'] or u['assegnato'] != u['mur']:
            raise ValueError(f'Importo incoerente: {u["codice"]}')
        if not u['ente'] or len(u['cf']) != 11:
            raise ValueError(f'Ente incompleto: {u["codice"]}')
        by_project[u['progetto']].append(u)
    if set(by_project) != {p['codice'] for p in projects}:
        raise ValueError('Unita fuori dai progetti elencati')
    for p, (cof, mur, assigned) in zip(projects, totals):
        group = sorted(by_project[p['codice']], key=lambda u: u['n'])
        if sum(u['ruolo'] == 'PI' for u in group) != 1 or [u['n'] for u in group] != list(range(1, len(group)+1)):
            raise ValueError(f'{p["codice"]}: numerazione ruoli non coerente')
        if (sum(u['cof'] for u in group), sum(u['mur'] for u in group), mur) != (cof, mur, assigned):
            raise ValueError(f'{p["codice"]}: subtotali discordanti')
        p.update({'cof': cof, 'mur': mur, 'n_unita': len(group)})
    if sum(p['mur'] for p in projects) != grand_total:
        raise ValueError(f'{sector}: totale finale non coerente')
    units.sort(key=lambda u: (next(p['n'] for p in projects if p['codice'] == u['progetto']), u['n']))
    return pdf, projects, units


def create_workbook(sector):
    pdf, projects, units = parse_pdf(sector)
    workbook = load_workbook(TEMPLATE)
    summary = workbook['Riepilogo']
    detail = workbook['Progetti']
    st_s = [copy(summary.cell(5, j)._style) for j in range(1, 7)]
    st_d = [copy(detail.cell(5, j)._style) for j in range(1, 14)]
    summary.delete_rows(5, summary.max_row - 4)
    detail.delete_rows(5, detail.max_row - 4)
    summary['A1'] = f'{sector} - Tabella B - Progetti finanziati'
    detail['A1'] = f'{sector} - Tabella B - Progetti finanziati - Dettaglio delle unità'
    detail['A2'] = 'Dati estratti dal PDF; in verde i progetti con PI under 40'
    for row_n, p in enumerate(projects, 5):
        values = [p['n'], p['codice'], p['titolo'], p['n_unita'], p['mur'], p['cof']]
        for j, value in enumerate(values, 1):
            cell = summary.cell(row_n, j, value)
            cell._style = copy(st_s[j-1])
    total_row = len(projects) + 5
    for j, val in [(1, 'TOTALE'), (4, len(units)), (5, sum(p['mur'] for p in projects)),
                   (6, sum(p['cof'] for p in projects))]:
        cell = summary.cell(total_row, j, val)
        cell._style = copy(st_s[j-1])
        font = copy(cell.font)
        font.bold = True
        cell.font = font
    summary.cell(total_row+2, 1, 'Fonte: ' + pdf.name)

    proj_by_code = {p['codice']: p for p in projects}
    green = SETTINGS[sector]['green']
    for row_n, u in enumerate(units, 5):
        p = proj_by_code[u['progetto']]
        under40 = 'Evidenziato in verde' if p['codice'] in green else 'Non evidenziato'
        values = [p['n'], u['progetto'], p['titolo'], u['ruolo'], u['responsabile'],
                  u['codice'], u['ente'], u['cf'], u['cof'], u['mur'], u['costo'],
                  u['assegnato'], under40]
        for j, value in enumerate(values, 1):
            cell = detail.cell(row_n, j, value)
            cell._style = copy(st_d[j-1])
        detail.cell(row_n, 8).number_format = '@'
        if p['codice'] in green:
            detail.cell(row_n, 13).fill = PatternFill('solid', fgColor='E2F0D9')
    detail.auto_filter.ref = f'A4:M{len(units)+4}'
    summary.auto_filter.ref = f'A4:F{len(projects)+4}'
    output = ROOT / 'data' / f'{sector}_Tabella_B_Progetti_finanziati.xlsx'
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)
    workbook.close()
    print(f'{output}: {len(projects)} progetti, {len(units)} unita, '
          f'{len(units)-len(projects)} RU; contributo MUR {sum(p["mur"] for p in projects):,} EUR')
    return output


def main():
    sectors = [sys.argv[1].upper()] if len(sys.argv) > 1 and sys.argv[1].upper() != 'ALL' else ['HPC', 'QT']
    for s in sectors:
        if s not in SETTINGS:
            raise SystemExit('Settore non valido: selezionare HPC, QT oppure ALL')
        create_workbook(s)


if __name__ == '__main__':
    main()
