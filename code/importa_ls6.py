#!/usr/bin/env python3
"""Importa LS6 dal PDF nella struttura Excel PRIN, validando le somme.

Uso: python code/importa_ls6.py [file.pdf]
Dipendenze: openpyxl, pdftotext (Poppler).
"""
from __future__ import annotations

import re
import subprocess
import sys
from collections import defaultdict
from copy import copy
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import PatternFill

ROOT = Path(__file__).resolve().parent.parent
SOURCE = ROOT / 'LS6_Tabella_B_Progetti finanziati.pdf'
OUTPUT = ROOT / 'data' / 'LS6_Tabella_B_Progetti_finanziati.xlsx'
TEMPLATE = ROOT / 'data' / 'SH6_Tabella_B_Progetti_finanziati.xlsx'
AMOUNT = r'(\d[\d.]*,\d{2})'
PROJECT_RE = re.compile(r'^\s*(\d+)\.\s+(2026[A-Z0-9]+)\s+-\s+(.+?)\s*$')
UNIT_START_RE = re.compile(r'^\s*\d+\s+(?:PI|RU)\s+.*?2026[A-Z0-9]+_\d{3}\b')
UNIT_RE = re.compile(
    r'^\s*(\d+)\s+(PI|RU)\s+(.*?)\s+(2026[A-Z0-9]+_\d{3})\s+'
    r'(.*?)\s*(\d{11})\s+' + AMOUNT + r'\s+' + AMOUNT + r'\s+' + AMOUNT + r'\s+' + AMOUNT + r'\s*$'
)
TOTAL_RE = re.compile(r'^\s*Totale\s+' + AMOUNT + r'\s+' + AMOUNT + r'\s+' + AMOUNT + r'\s*$')
GRAND_RE = re.compile(r'^.*TOTALE COMPLESSIVO\s+' + AMOUNT + r'\s*$')


def euro_int(v: str) -> int:
    amount = Decimal(v.replace('.', '').replace(',', '.'))
    if amount != amount.to_integral_value():
        raise ValueError(f'Importo non intero: {v}')
    return int(amount)


def importa(pdf=SOURCE):
    pdf = Path(pdf)
    if not pdf.is_file():
        raise FileNotFoundError(pdf)
    if not TEMPLATE.is_file():
        raise FileNotFoundError(TEMPLATE)
    raw = subprocess.check_output(['pdftotext', '-layout', str(pdf), '-'], text=True)
    projects, units, totals = [], [], []
    grand_total = None
    pending_unit = None

    def parse_unit(line):
        match = UNIT_RE.match(' '.join(line.split()))
        if not match:
            raise ValueError(f'Impossibile estrarre riga unità dal PDF: {line}')
        n, role, name, unit_id, institution, cf, cofin, mur, cost, granted = match.groups()
        return {
            'n': int(n), 'ruolo': role, 'responsabile': name.strip(), 'codice': unit_id,
            'progetto': unit_id.split('_')[0], 'ente': institution.strip(), 'cf': cf,
            'cof': euro_int(cofin), 'mur': euro_int(mur),
            'costo': euro_int(cost), 'assegnato': euro_int(granted),
        }

    for line in raw.splitlines():
        if pending_unit is not None:
            pending_unit += ' ' + line.strip()
            if UNIT_RE.match(' '.join(pending_unit.split())):
                units.append(parse_unit(pending_unit))
                pending_unit = None
            continue
        if match := PROJECT_RE.match(line):
            n, code, title = match.groups()
            projects.append({'n': int(n), 'codice': code, 'titolo': title.strip()})
        elif UNIT_START_RE.match(line):
            if UNIT_RE.match(' '.join(line.split())):
                units.append(parse_unit(line))
            else:
                pending_unit = line.strip()
        elif match := TOTAL_RE.match(line):
            totals.append(tuple(euro_int(x) for x in match.groups()))
        elif match := GRAND_RE.match(line):
            grand_total = euro_int(match.group(1))

    if pending_unit:
        raise ValueError('Ultima unità incompleta: ' + pending_unit)
    if len(projects) != 5 or len(units) != 25 or len(totals) != 5:
        raise ValueError(f'Estrazione incompleta: {len(projects)} progetti, {len(units)} unità, {len(totals)} totali')
    if grand_total != 5955206:
        raise ValueError(f'Totale inatteso rispetto al PDF: {grand_total}')
    by_project = defaultdict(list)
    if len({u['codice'] for u in units}) != len(units):
        raise ValueError('Codici unità ripetuti')
    for unit in units:
        if unit['cof'] + unit['mur'] != unit['costo'] or unit['mur'] != unit['assegnato']:
            raise ValueError('Importi discordanti per ' + unit['codice'])
        by_project[unit['progetto']].append(unit)
    for p, (cof, mur, granted) in zip(projects, totals):
        group = by_project[p['codice']]
        if not group or sum(u['ruolo'] == 'PI' for u in group) != 1:
            raise ValueError(f'Numero PI sbagliato per {p["codice"]}')
        if [u['n'] for u in group] != list(range(1, len(group) + 1)):
            raise ValueError(f'Sequenza numerica errata per {p["codice"]}')
        if sum(u['cof'] for u in group) != cof or sum(u['mur'] for u in group) != mur or granted != mur:
            raise ValueError(f'Totale progetto discordante {p["codice"]}')
        p.update({'cof': cof, 'mur': mur, 'n_unita': len(group)})
    if sum(p['mur'] for p in projects) != grand_total:
        raise ValueError('Totale generale discordante')

    wb = load_workbook(TEMPLATE)
    summary, detail = wb['Riepilogo'], wb['Progetti']
    summary_styles = [copy(summary.cell(5, j)._style) for j in range(1, 7)]
    detail_styles = [copy(detail.cell(5, j)._style) for j in range(1, 14)]
    summary.delete_rows(5, summary.max_row - 4)
    detail.delete_rows(5, detail.max_row - 4)
    summary['A1'] = 'LS6 - Tabella B - Progetti finanziati'
    detail['A1'] = 'LS6 - Tabella B - Progetti finanziati - Dettaglio delle unità'
    detail['A2'] = 'Dati estratti dal PDF; in verde i progetti con PI under 40'
    for i, p in enumerate(projects, 5):
        for j, value in enumerate([p['n'], p['codice'], p['titolo'], p['n_unita'], p['mur'], p['cof']], 1):
            cell = summary.cell(i, j, value)
            cell._style = copy(summary_styles[j - 1])
    totals_row = len(projects) + 5
    for j, value in [(1, 'TOTALE'), (4, len(units)),
                     (5, sum(p['mur'] for p in projects)), (6, sum(p['cof'] for p in projects))]:
        cell = summary.cell(totals_row, j, value)
        cell._style = copy(summary_styles[j - 1])
        font = copy(cell.font)
        font.bold = True
        cell.font = font
    summary.cell(totals_row + 2, 1, f'Fonte: {pdf.name}')
    by_code = {p['codice']: p for p in projects}
    green_project = '2026PA8574'  # Evidenziato in verde nella pagina PDF.
    for i, u in enumerate(units, 5):
        p = by_code[u['progetto']]
        under40 = 'Evidenziato in verde' if u['progetto'] == green_project else 'Non evidenziato'
        row = [p['n'], u['progetto'], p['titolo'], u['ruolo'], u['responsabile'],
               u['codice'], u['ente'], u['cf'], u['cof'], u['mur'], u['costo'], u['assegnato'], under40]
        for j, value in enumerate(row, 1):
            cell = detail.cell(i, j, value)
            cell._style = copy(detail_styles[j - 1])
        detail.cell(i, 8).number_format = '@'
        if u['progetto'] == green_project:
            detail.cell(i, 13).fill = PatternFill('solid', fgColor='E2F0D9')
    detail.auto_filter.ref = f'A4:M{len(units)+4}'
    summary.auto_filter.ref = f'A4:F{len(projects)+4}'
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUTPUT)
    wb.close()
    print(f'{OUTPUT}: {len(projects)} progetti, {len(units)} unità, '
          f'{len(units)-len(projects)} RU, MUR {grand_total:,.0f} EUR')
    return OUTPUT


if __name__ == '__main__':
    importa(Path(sys.argv[1]) if len(sys.argv) > 1 else SOURCE)
