"""Lettura e validazione dei file PRIN 2026 (due formati Excel supportati).

Schema A: foglio Progetti = unita; Riepilogo/Totali progetti = importi progetto.
Schema B: foglio Unita = unita; Progetti = importi progetto.
"""
from collections import defaultdict
from dataclasses import dataclass
from decimal import Decimal, InvalidOperation
from pathlib import Path
import re
import unicodedata

from openpyxl import load_workbook

# Percorsi stabili rispetto alla posizione dei sorgenti, non alla working directory.
PROJECT_ROOT = Path(__file__).resolve().parent.parent
DATA_DIR = PROJECT_ROOT / 'data'
CODICE_PROGETTO = re.compile(r'^20\d{2}[A-Z0-9]+$')
SETTORE = re.compile(r'^(?:PE\d+|LS\d+|SH\d+|CYB|IA)$')


def normalizza(s):
    s = unicodedata.normalize('NFKD', str(s or '').casefold())
    return ''.join(c for c in s if c.isalnum())


def importo(valore):
    if valore is None or str(valore).strip() == '':
        raise ValueError('Importo MUR assente')
    if isinstance(valore, (int, float, Decimal)):
        return Decimal(str(valore))
    s = str(valore).replace('€', '').replace('\u00a0', '').replace(' ', '').strip()
    if ',' in s:
        s = s.replace('.', '').replace(',', '.')
    try:
        return Decimal(s)
    except InvalidOperation as exc:
        raise ValueError(f'Importo MUR non numerico: {valore!r}') from exc


def euro(x):
    return f'{x:,.2f}'.replace(',', 'X').replace('.', ',').replace('X', '.')


def intestazioni(foglio, obbligatorie):
    """Trova l'intestazione anche dopo righe di titolo/note introduttive."""
    for num, row in enumerate(foglio.iter_rows(values_only=True), 1):
        colonne = {normalizza(v): i for i, v in enumerate(row) if v is not None}
        if all(any(c == req or c.startswith(req) for c in colonne) for req in obbligatorie):
            return num, colonne
    raise ValueError(f'{foglio.title}: intestazioni non trovate: {obbligatorie}')


def colonna(headers, *possibili, required=True):
    for s in possibili:
        n = normalizza(s)
        if n in headers:
            return headers[n]
    for s in possibili:
        n = normalizza(s)
        for k, idx in headers.items():
            if k.startswith(n):
                return idx
    if required:
        raise ValueError(f'Colonna assente: {possibili}; presenti: {list(headers)}')
    return None


def cella(row, idx):
    return row[idx] if idx is not None and idx < len(row) else None


@dataclass(frozen=True)
class Unita:
    settore: str
    progetto: str
    ruolo: str
    responsabile: str
    codice_unita: str
    ente: str
    cf: str
    mur: Decimal


@dataclass(frozen=True)
class Progetto:
    settore: str
    codice: str
    mur: Decimal


@dataclass
class Dataset:
    file: list
    unita: list
    progetti: dict
    avvisi: list


def carica_dataset(cartella=DATA_DIR):
    cartella = Path(cartella)
    if not cartella.is_dir():
        raise FileNotFoundError(f'Cartella dati inesistente: {cartella}')
    file = sorted(f for f in cartella.glob('*.xlsx') if not f.name.startswith('~$'))
    if not file:
        raise ValueError(f'Non ci sono file .xlsx in {cartella}')

    unita, progetti, avvisi = [], {}, []
    for percorso in file:
        settore = percorso.name.split('_', 1)[0].upper()
        if not SETTORE.fullmatch(settore):
            raise ValueError(f'{percorso.name}: settore non riconosciuto: {settore}')
        wb = load_workbook(percorso, read_only=True, data_only=True)
        try:
            # Due convenzioni adottate nei file Excel caricati.
            if 'Unita' in wb.sheetnames:
                dettaglio, riepilogo = wb['Unita'], wb['Progetti']
            else:
                dettaglio = wb['Progetti']
                riepilogo = wb['Totali progetti'] if 'Totali progetti' in wb.sheetnames else wb['Riepilogo']

            rsum, hs = intestazioni(riepilogo, ('codiceprogetto',))
            ip = colonna(hs, 'Codice progetto')
            im = colonna(hs, 'Contributo MUR assegnato', 'Contributo MUR', 'MUR assegnato')
            iu_att = colonna(hs, 'N. unità', 'N° unità', required=False)
            iru_att = colonna(hs, 'N. RU', required=False)
            attese = {}
            for riga in riepilogo.iter_rows(min_row=rsum + 1, values_only=True):
                codice = str(cella(riga, ip) or '').strip().upper()
                if not CODICE_PROGETTO.fullmatch(codice):
                    continue  # totali, note e righe vuote
                key = (settore, codice)
                if key in progetti:
                    raise ValueError(f'{percorso.name}: codice progetto duplicato: {codice}')
                try:
                    valore = importo(cella(riga, im))
                except ValueError as e:
                    raise ValueError(f'{percorso.name}: progetto {codice}: {e}') from e
                progetti[key] = Progetto(settore, codice, valore)
                attese[key] = (cella(riga, iu_att), cella(riga, iru_att))

            rd, hd = intestazioni(dettaglio, ('codiceprogetto', 'ruolo', 'ateneoente'))
            idp = colonna(hd, 'Codice progetto')
            idr = colonna(hd, 'Ruolo')
            idnome = colonna(hd, 'Responsabile unità', 'Responsabile')
            idun = colonna(hd, 'Codice unità')
            idente = colonna(hd, 'Ateneo/Ente')
            idcf = colonna(hd, 'Codice fiscale', required=False)
            idmur = colonna(hd, 'Contributo MUR assegnato', 'MUR assegnato')
            conteggi = defaultdict(lambda: {'PI': 0, 'RU': 0, 'mur': Decimal(0)})
            codici_unita = set()
            for num, riga in enumerate(dettaglio.iter_rows(min_row=rd + 1, values_only=True), rd + 1):
                codice = str(cella(riga, idp) or '').strip().upper()
                if not CODICE_PROGETTO.fullmatch(codice):
                    continue
                ruolo = str(cella(riga, idr) or '').strip().upper()
                if ruolo not in ('PI', 'RU'):
                    raise ValueError(f'{percorso.name}:{num}: ruolo non valido {ruolo!r}')
                persona = ' '.join(str(cella(riga, idnome) or '').split())
                ente = ' '.join(str(cella(riga, idente) or '').split())
                cf = str(cella(riga, idcf) or '').strip()
                codice_unita = str(cella(riga, idun) or '').strip()
                if not persona or not ente or not codice_unita:
                    raise ValueError(f'{percorso.name}:{num}: responsabile, ente o codice unità mancanti')
                try:
                    mur = importo(cella(riga, idmur))
                except ValueError as e:
                    raise ValueError(f'{percorso.name}:{num}: {e}') from e
                key = (settore, codice)
                if key not in progetti:
                    raise ValueError(f'{percorso.name}:{num}: progetto {codice} non presente nel riepilogo')
                if codice_unita in codici_unita:
                    avvisi.append(f'{percorso.name}: codice unità ripetuto {codice_unita} '
                                  f'(conservate entrambe le righe, senza perdere i finanziamenti)')
                codici_unita.add(codice_unita)
                unita.append(Unita(settore, codice, ruolo, persona, codice_unita, ente, cf, mur))
                conteggi[key][ruolo] += 1
                conteggi[key]['mur'] += mur

            for key, (n_unita, n_ru) in attese.items():
                c = conteggi[key]
                if not c['PI'] == 1:
                    raise ValueError(f'{percorso.name}: {key[1]}: atteso 1 PI, trovati {c["PI"]}')
                if n_unita is not None and c['PI'] + c['RU'] != int(n_unita):
                    avvisi.append(f'{percorso.name}: {key[1]}: nel riepilogo N. unità={n_unita}, '
                                  f'ma righe di dettaglio={c["PI"] + c["RU"]}')
                if n_ru is not None and c['RU'] != int(n_ru):
                    avvisi.append(f'{percorso.name}: {key[1]}: nel riepilogo N. RU={n_ru}, '
                                  f'ma righe RU di dettaglio={c["RU"]}')
                if c['mur'] != progetti[key].mur:
                    raise ValueError(f'{percorso.name}: {key[1]}: MUR unità {euro(c["mur"])} '
                                     f'!= MUR riepilogo {euro(progetti[key].mur)}')
        finally:
            wb.close()
    return Dataset(file, unita, progetti, avvisi)


def nome_ente(insieme):
    """Preferisce una descrizione breve e stabile fra le varianti di uno stesso CF."""
    return sorted(insieme, key=lambda x: (len(x), x.casefold()))[0]


def chiave_ente(u):
    return ('cf', u.cf) if u.cf else ('nome', normalizza(u.ente))
