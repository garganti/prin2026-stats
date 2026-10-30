#!/usr/bin/env python3
"""Importa il PDF PRIN SH7 nello schema Excel Riepilogo + Progetti del progetto.

Uso: python code/importa_sh7.py [percorso_del_pdf]
Dipendenze: openpyxl e pdftotext (Poppler).
"""
from __future__ import annotations

import re
import subprocess
from collections import defaultdict
from copy import copy
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import PatternFill

ROOT = Path(__file__).resolve().parent.parent
SOURCE = ROOT / 'SH7_Tabella_B_Progetti finanziati.pdf'
OUTPUT = ROOT / 'data' / 'SH7_Tabella_B_Progetti_finanziati.xlsx'
TEMPLATE = ROOT / 'data' / 'SH6_Tabella_B_Progetti_finanziati.xlsx'
AMOUNT = r'(\d[\d.]*,\d{2})'
PROJECT_RE = re.compile(r'^\s*(\d+)\.\s+(2026[A-Z0-9]+)\s+-\s+(.+?)\s*$')
UNIT_RE = re.compile(
    r'^\s*(\d+)\s+(PI|RU)\s+(.*?)\s+(2026[A-Z0-9]+_\d{3})\s+'
    r'(.*?)\s+(\d{11})\s+' + AMOUNT + r'\s+' + AMOUNT + r'\s+' + AMOUNT + r'\s+' + AMOUNT + r'\s*$'
)
TOTAL_RE = re.compile(r'^\s*Totale\s+' + AMOUNT + r'\s+' + AMOUNT + r'\s+' + AMOUNT + r'\s*$')
GRAND_RE = re.compile(r'^.*TOTALE COMPLESSIVO\s+' + AMOUNT + r'\s*$')

# In PDF alcune ragioni sociali occupano piu' righe, separate dalla riga della persona.
MULTILINE_ENTITIES = {
    '2026P53P8J_001': 'Gran Sasso Science Institute - Scuola di dottorato internazionale',
    '2026PXHN45_003': "Scuola Superiore di Studi Universitari e Perfezionamento Sant'Anna",
}


def integer_euro(v):
    value = v.replace('.', '').replace(',', '.')
    number = Decimal(value)
    if number != number.to_integral_value():
        raise ValueError(f'Si attende un importo intero in euro: {v}')
    return int(number)


def importa(pdf=SOURCE):
    pdf = Path(pdf)
    if not pdf.is_file() or not TEMPLATE.is_file():
        raise FileNotFoundError(f'PDF o file modello mancante: {pdf}, {TEMPLATE}')
    raw = subprocess.check_output(['pdftotext', '-layout', str(pdf), '-'], text=True)
    progetti = []
    unita = []
    totali = []
    totale_generale = None
    for line in raw.splitlines():
        if match := PROJECT_RE.match(line):
            n, codice, titolo = match.groups()
            progetti.append({'n': int(n), 'codice': codice, 'titolo': titolo})
            continue
        if match := UNIT_RE.match(line):
            n, ruolo, responsabile, codice, ente, cf, cof, mur, costo, assegnato = match.groups()
            codice_progetto = codice.split('_', 1)[0]
            ente = MULTILINE_ENTITIES.get(codice, ente.strip())
            unita.append({
                'n': int(n), 'ruolo': ruolo, 'responsabile': ' '.join(responsabile.split()),
                'codice_unita': codice, 'progetto': codice_progetto, 'ente': ' '.join(ente.split()),
                'cf': cf, 'cof': integer_euro(cof), 'mur': integer_euro(mur),
                'costo': integer_euro(costo), 'assegnato': integer_euro(assegnato),
            })
            continue
        if match := TOTAL_RE.match(line):
            cof, mur, assegnato = match.groups()
            totali.append((integer_euro(cof), integer_euro(mur), integer_euro(assegnato)))
            continue
        if match := GRAND_RE.match(line):
            totale_generale = integer_euro(match.group(1))

    if len(progetti) != 6 or len(unita) != 31 or len(totali) != 6:
        raise ValueError(f'Estrazione incompleta: {len(progetti)} progetti, '
                         f'{len(unita)} unita, {len(totali)} totali')
    if totale_generale != 7001948:
        raise ValueError(f'Totale nel PDF inatteso: {totale_generale}')
    codes = {p['codice'] for p in progetti}
    if len(codes) != len(progetti) or len({u['codice_unita'] for u in unita}) != len(unita):
        raise ValueError('Codici progetti o unita duplicati')
    byproject = defaultdict(list)
    for u in unita:
        if u['progetto'] not in codes:
            raise ValueError(f"Progetto sconosciuto: {u['progetto']}")
        if not u['ente'] or not u['responsabile'] or len(u['cf']) != 11:
            raise ValueError(f"Campi mancanti: {u['codice_unita']}")
        if u['cof'] + u['mur'] != u['costo'] or u['mur'] != u['assegnato']:
            raise ValueError(f"Importi unita discordanti: {u['codice_unita']}")
        byproject[u['progetto']].append(u)
    for p, (cof, mur, assegnato) in zip(progetti, totali):
        v = byproject[p['codice']]
        if len([u for u in v if u['ruolo'] == 'PI']) != 1:
            raise ValueError(f"PI inatteso: {p['codice']}")
        if sum(u['cof'] for u in v) != cof or sum(u['mur'] for u in v) != mur or mur != assegnato:
            raise ValueError(f"Totali non coincidenti nel progetto {p['codice']}")
        if [u['n'] for u in v] != list(range(1, len(v)+1)):
            raise ValueError(f"Sequenza unita incompleta: {p['codice']}")
        p.update({'cof': cof, 'mur': mur, 'n_unita': len(v)})
    if sum(p['mur'] for p in progetti) != totale_generale:
        raise ValueError('Totale complessivo non corrispondente al PDF')

    wb = load_workbook(TEMPLATE)
    sh_r = wb['Riepilogo']
    sh_u = wb['Progetti']
    styles_r = [copy(sh_r.cell(5, c)._style) for c in range(1, 7)]
    styles_u = [copy(sh_u.cell(5, c)._style) for c in range(1, 14)]
    sh_r.delete_rows(5, sh_r.max_row - 4)
    sh_u.delete_rows(5, sh_u.max_row - 4)
    sh_r['A1'] = 'SH7 - Tabella B - Progetti finanziati'
    sh_u['A1'] = 'SH7 - Tabella B - Progetti finanziati - Dettaglio delle unità'
    sh_u['A2'] = 'Dati estratti dal PDF; in verde i progetti con PI under 40'

    for i, p in enumerate(progetti, 5):
        row = [p['n'], p['codice'], p['titolo'], p['n_unita'], p['mur'], p['cof']]
        for j, value in enumerate(row, 1):
            cell = sh_r.cell(i, j, value)
            cell._style = copy(styles_r[j-1])
    total_row = len(progetti) + 5
    for j, value in [(1, 'TOTALE'), (4, len(unita)),
                     (5, sum(p['mur'] for p in progetti)), (6, sum(p['cof'] for p in progetti))]:
        cell = sh_r.cell(total_row, j, value)
        cell._style = copy(styles_r[j-1])
        new_font = copy(cell.font)
        new_font.bold = True
        cell.font = new_font
    sh_r.cell(total_row+2, 1, f'Fonte: {pdf.name}')

    project_map = {p['codice']: p for p in progetti}
    # Titoli con caratteri verdi nel documento sono FAMBIT e INFLOW.
    under40_codes = {'2026PFPY3A', '2026PNA9YP'}
    for i, u in enumerate(unita, 5):
        p = project_map[u['progetto']]
        status = 'Evidenziato in verde' if u['progetto'] in under40_codes else 'Non evidenziato'
        row = [p['n'], u['progetto'], p['titolo'], u['ruolo'], u['responsabile'],
               u['codice_unita'], u['ente'], u['cf'], u['cof'], u['mur'], u['costo'],
               u['assegnato'], status]
        for j, value in enumerate(row, 1):
            cell = sh_u.cell(i, j, value)
            cell._style = copy(styles_u[j-1])
        sh_u.cell(i, 8).number_format = '@'
        if p['codice'] in under40_codes:
            sh_u.cell(i, 13).fill = PatternFill('solid', fgColor='E2F0D9')
    sh_u.auto_filter.ref = f'A4:M{len(unita)+4}'
    sh_r.auto_filter.ref = f'A4:F{len(progetti)+4}'
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUTPUT)
    print(f'{OUTPUT}: {len(progetti)} progetti, {len(unita)} unita, '
          f'{len(unita) - len(progetti)} RU, MUR {totale_generale:,.0f} EUR')
    return OUTPUT


if __name__ == '__main__':
    import sys
    importa(Path(sys.argv[1]) if len(sys.argv) > 1 else SOURCE)
