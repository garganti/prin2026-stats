from collections import defaultdict
from pathlib import Path
import re
import unicodedata
from decimal import Decimal
from openpyxl import load_workbook

DATA_DIR = Path(__file__).resolve().parent / 'data'
OUTPUT = Path(__file__).resolve().parent / 'riepilogo_settori.md'


def normalizza(testo):
    testo = unicodedata.normalize('NFKD', str(testo).casefold())
    return ''.join(c for c in testo if c.isalnum())


def righe_con_intestazioni(foglio, richieste):
    for numero, riga in enumerate(foglio.iter_rows(values_only=True), 1):
        intestazioni = {normalizza(valore): i for i, valore in enumerate(riga) if valore is not None}
        if all(any(nome in chiave for chiave in intestazioni) for nome in richieste):
            return numero, intestazioni
    raise ValueError(f'Intestazioni non trovate nel foglio {foglio.title}')


def colonna(intestazioni, *alternative):
    for alternativa in alternative:
        chiave = normalizza(alternativa)
        for nome, indice in intestazioni.items():
            if nome == chiave or nome.startswith(chiave):
                return indice
    raise ValueError(f'Colonna non trovata: {alternative}')


def numero(valore):
    if isinstance(valore, (int, float, Decimal)):
        return Decimal(str(valore))
    if valore is None:
        return Decimal(0)
    valore = str(valore).strip().replace('€', '').replace(' ', '')
    if ',' in valore:
        valore = valore.replace('.', '').replace(',', '.')
    return Decimal(valore)


def euro(valore):
    return f'{valore:,.2f}'.replace(',', 'X').replace('.', ',').replace('X', '.')


riepilogo = defaultdict(lambda: {'progetti': set(), 'ru': set(), 'importi': {}})
files = sorted(DATA_DIR.glob('*.xlsx'))
if not files:
    raise SystemExit(f'Nessun Excel in {DATA_DIR}')

for percorso in files:
    settore = percorso.name.split('_', 1)[0].upper()
    if not re.fullmatch(r'(PE|LS|SH)\d+|CYB', settore):
        print(f'Salto il file senza settore riconosciuto: {percorso.name}')
        continue
    workbook = load_workbook(percorso, read_only=True, data_only=True)
    try:
        dettaglio = workbook['Progetti']
        riga_header, headers = righe_con_intestazioni(dettaglio, ['codiceprogetto', 'ruolo'])
        c_progetto = colonna(headers, 'Codice progetto')
        c_ruolo = colonna(headers, 'Ruolo')
        c_unita = colonna(headers, 'Codice unità')

        for riga in dettaglio.iter_rows(min_row=riga_header + 1, values_only=True):
            codice = riga[c_progetto]
            ruolo = str(riga[c_ruolo] or '').strip().upper()
            if codice and ruolo in ('PI', 'RU'):
                riepilogo[settore]['progetti'].add(str(codice).strip())
                if ruolo == 'RU':
                    # Ogni unita di ricerca viene contata una sola volta.
                    identificatore = str(riga[c_unita] or '').strip() or f'{codice}:{len(riepilogo[settore]["ru"])}'
                    riepilogo[settore]['ru'].add(identificatore)

        nome_foglio = 'Totali progetti' if 'Totali progetti' in workbook else 'Riepilogo'
        foglio_importi = workbook[nome_foglio]
        riga_header, headers = righe_con_intestazioni(foglio_importi, ['codiceprogetto', 'contributomur'])
        c_progetto = colonna(headers, 'Codice progetto')
        c_importo = colonna(headers, 'Contributo MUR assegnato', 'Contributo MUR')
        for riga in foglio_importi.iter_rows(min_row=riga_header + 1, values_only=True):
            codice = riga[c_progetto]
            if not isinstance(codice, str) or not re.fullmatch(r'2026[A-Z0-9]+', codice.strip()):
                continue  # Esclude righe Totale e note
            importo = numero(riga[c_importo])
            vecchio = riepilogo[settore]['importi'].get(codice)
            if vecchio is not None and vecchio != importo:
                raise ValueError(f'Importi discordanti per {codice}')
            riepilogo[settore]['importi'][codice] = importo
    finally:
        workbook.close()

righe = ['# Riepilogo progetti finanziati per settore', '',
         '| Settore | Progetti | RU | Contributo MUR totale (€) |',
         '|:---|---:|---:|---:|']
tot_progetti = tot_ru = 0
tot_euro = Decimal(0)
for settore, dati in sorted(riepilogo.items()):
    senza_importo = dati['progetti'] - dati['importi'].keys()
    if senza_importo:
        raise ValueError(f'{settore}: progetti privi di importo: {sorted(senza_importo)}')
    n_progetti = len(dati['progetti'])
    n_ru = len(dati['ru'])
    importo = sum((dati['importi'][c] for c in dati['progetti']), Decimal(0))
    righe.append(f'| {settore} | {n_progetti} | {n_ru} | {euro(importo)} |')
    tot_progetti += n_progetti
    tot_ru += n_ru
    tot_euro += importo
righe.append(f'| **TOTALE** | **{tot_progetti}** | **{tot_ru}** | **{euro(tot_euro)}** |')
righe.extend(['', 'Gli RU rappresentano le unità di ricerca con ruolo RU (esclusi i PI).',
              'Il contributo totale è la somma degli importi MUR per progetto, senza duplicazioni.'])
OUTPUT.write_text('\n'.join(righe) + '\n', encoding='utf-8')
print('\n'.join(righe))
print(f'\nSalvato: {OUTPUT}')
